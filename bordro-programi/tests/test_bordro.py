"""Testler: python -m pytest tests"""
import csv
import os
import sys
from decimal import Decimal

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from bordro import iceaktar, raporlar  # noqa: E402
from bordro.hesap import (D, EkKalem, PARAMETRE_TANIMLARI, Personel, VARSAYILAN_DILIMLER,  # noqa: E402
                          hesapla, kumulatif_vergi, vergi_tarifesi)
from bordro.veritabani import DonemKilitli, Veritabani, iban_gecerli, tc_gecerli  # noqa: E402

PARAM = {k: v for k, (_, v) in PARAMETRE_TANIMLARI.items()}
DILIM = [(None if u is None else D(u), D(o)) for u, o in VARSAYILAN_DILIMLER[2026]]


def hemsire(**kw):
    v = dict(tc="12345678950", ad="Ayşe", soyad="Yılmaz", unvan="Hemşire", birim="Acil", derece=5, kademe=3,
             hizmet_yili=12, yan_odeme_puani=2200, ozel_hizmet_orani=D(100), ek_odeme_orani=D(87),
             es_yardimi=True, cocuk_0_6=1, iban="TR330006100519786457841326")
    v.update(kw)
    return Personel(**v)


def test_vergi_tarifesi():
    assert vergi_tarifesi(D(100000), DILIM) == D("15000")
    # 190.000 x %15 + 10.000 x %20
    assert vergi_tarifesi(D(200000), DILIM) == D("30500")
    # dilim geçişi: 180.000'den sonra 20.000 matrah -> 10.000 x %15 + 10.000 x %20
    assert kumulatif_vergi(D(180000), D(20000), DILIM) == D("3500")
    assert vergi_tarifesi(D(0), DILIM) == 0


def test_657_memur_elle_dogrulanmis():
    """Değerler hesap makinesiyle tek tek doğrulandı (Temmuz 2026 katsayıları)."""
    s = hesapla(hemsire(), "2026-07", PARAM, DILIM, gosterge=895)
    o = {x.ad: x.tutar for x in s.odemeler}
    assert o["Aylık (Gösterge)"] == D("1410.08")          # 895 x 1,575512
    assert o["Taban Aylık"] == D("25794.92")              # 1000 x 25,794915
    assert o["Kıdem Aylığı"] == D("378.12")               # 12 x 20 x 1,575512
    assert o["Yan Ödeme"] == D("1099.23")                 # 2200 x 0,499649
    assert o["Özel Hizmet Tazminatı"] == D("40762.28")    # (9500 x AK + 1000 x TK) x %100
    assert o["Ek Ödeme (666 KHK)"] == D("35463.18")       # aynı taban x %87
    assert o["Aile Yardımı (Eş)"] == D("3581.14")         # 2273 x AK
    assert o["Çocuk Yardımı"] == D("787.76")              # 500 x AK
    assert s.sgk_matrahi == D("104907.81")                # aile/çocuk hariç
    k = {x.ad: x.tutar for x in s.kesintiler}
    assert k["SGK Malullük-Yaşlılık (Kişi)"] == D("9441.70")
    assert k["GSS Primi (Kişi)"] == D("5245.39")
    assert s.gv_matrahi == D("90220.72")
    # 13.533,11 hesaplanan GV - 4.537,75 asgari ücret istisnası (Temmuz)
    assert k["Gelir Vergisi"] == D("8995.36")
    # 104.907,81 x %0,759 = 796,25 - 250,70 istisna
    assert k["Damga Vergisi"] == D("545.55")
    assert s.net == s.brut - s.kesinti_toplami
    assert {x.ad for x in s.isveren} == {"SGK Malullük-Yaşlılık (Kurum)", "GSS Primi (Kurum)"}


def test_5434_emekli_keseneği_yalniz_aylik_kalemlerinden():
    s = hesapla(hemsire(emeklilik="5434"), "2026-07", PARAM, DILIM, gosterge=895)
    kesenek_matrahi = D("1410.08") + D("25794.92") + D("378.12")
    assert s.kesinti("Emekli Keseneği (Kişi)") == (kesenek_matrahi * 16 / 100).quantize(D("0.01"))
    assert s.kesinti("SGK Malullük-Yaşlılık (Kişi)") == 0


def test_4b_sozlesmeli():
    p = Personel(tc="32345678996", ad="Zeynep", soyad="Öztürk", statu="4B", sozlesme_ucreti=D(40000))
    s = hesapla(p, "2026-01", PARAM, DILIM)
    assert s.brut == D(40000)
    assert s.kesinti("SGK Malullük-Yaşlılık (Kişi)") == D("3600.00")   # %14 - %5 GSS
    assert s.kesinti("GSS Primi (Kişi)") == D("2000.00")
    assert s.kesinti("İşsizlik Sigortası (Kişi)") == D("400.00")
    assert s.gv_matrahi == D("34000.00")
    # Ocak: 34.000 x %15 = 5.100 - 28.075,50 x %15 = 4.211,33 istisna
    assert s.kesinti("Gelir Vergisi") == D("888.67")
    assert not any(x.ad == "Kıdem Aylığı" for x in s.odemeler)


def test_sgk_tavani():
    p = Personel(tc="32345678996", ad="A", soyad="B", statu="4B", sozlesme_ucreti=D(500000))
    s = hesapla(p, "2026-01", PARAM, DILIM)
    assert s.sgk_matrahi == D("33030") * D("7.5")


def test_ek_kalemler_ve_engellilik():
    ek = [EkKalem("Nöbet Ücreti", D(5000), "odeme", vergiye_tabi=True, damgaya_tabi=True, sgk_tabi=False),
          EkKalem("İcra", D(1000), "kesinti")]
    temel = hesapla(hemsire(), "2026-07", PARAM, DILIM, gosterge=895)
    s = hesapla(hemsire(), "2026-07", PARAM, DILIM, gosterge=895, ek_kalemler=ek)
    assert s.brut == temel.brut + 5000
    assert s.sgk_matrahi == temel.sgk_matrahi
    assert s.gv_matrahi == temel.gv_matrahi + 5000
    assert s.kesinti("İcra") == 1000
    param = {**PARAM, "engelli_indirimi_2": "6000"}
    e = hesapla(hemsire(engellilik=2), "2026-07", param, DILIM, gosterge=895)
    assert e.gv_matrahi == temel.gv_matrahi - 6000


def test_kumulatif_matrah_yuksek_dilime_tasir():
    ilk = hesapla(hemsire(), "2026-07", PARAM, DILIM, gosterge=895)
    sonra = hesapla(hemsire(), "2026-07", PARAM, DILIM, gosterge=895, onceki_kumulatif_matrah=D(1000000))
    assert sonra.kesinti("Gelir Vergisi") > ilk.kesinti("Gelir Vergisi")
    assert sonra.kumulatif_matrah == D(1000000) + sonra.gv_matrahi


def test_kimlik_ve_iban_dogrulama():
    assert tc_gecerli("12345678950")
    assert not tc_gecerli("12345678951")
    assert not tc_gecerli("02345678950")
    assert iban_gecerli("TR33 0006 1005 1978 6457 8413 26")
    assert not iban_gecerli("TR330006100519786457841327")


@pytest.fixture
def db():
    v = Veritabani(":memory:")
    yield v
    v.kapat()


def test_veritabani_dongusu(db, tmp_path):
    pid = db.personel_kaydet(hemsire())
    db.personel_kaydet(Personel(tc="32345678996", ad="Zeynep", soyad="Öztürk", birim="Doğumhane", statu="4B",
                                sozlesme_ucreti=D(40000)))
    db.ek_kalem_ekle(pid, "2026-07", EkKalem("İcra", D(750), "kesinti"))
    db.ek_kalem_ekle(pid, None, EkKalem("Lojman", D(100), "kesinti"))

    haziran = db.donem_hesapla("2026-06")
    assert haziran[0].odeme("Taban Aylık") == D("22722.79")  # Ocak 2026 seti geçerli
    db.bordro_kaydet("2026-06", haziran)
    temmuz = db.donem_hesapla("2026-07")
    ayse = next(s for s in temmuz if s.personel.id == pid)
    assert ayse.kesinti("İcra") == 750 and ayse.kesinti("Lojman") == 100
    onceki = next(s for s in haziran if s.personel.id == pid).gv_matrahi
    assert ayse.kumulatif_matrah == onceki + ayse.gv_matrahi

    db.bordro_kaydet("2026-07", temmuz)
    geri = db.kayitli_bordro("2026-07")
    assert [s.net for s in geri] == [s.net for s in sorted(temmuz, key=lambda s: (s.personel.birim,
                                                                                  s.personel.ad))]
    db.donem_kilitle("2026-07")
    with pytest.raises(DonemKilitli):
        db.bordro_kaydet("2026-07", temmuz)

    # bordrosu olan personel silinmez, pasife alınır
    db.personel_sil(pid)
    assert not db.personel(pid).aktif


def test_personel_dogrulama(db):
    with pytest.raises(ValueError):
        db.personel_kaydet(hemsire(tc="11111111111"))
    with pytest.raises(ValueError):
        db.personel_kaydet(hemsire(iban="TR000000000000000000000000"))
    with pytest.raises(ValueError):
        db.personel_kaydet(hemsire(derece=5, kademe=9))  # gösterge tablosunda yok
    db.personel_kaydet(hemsire())
    with pytest.raises(ValueError):
        db.personel_kaydet(hemsire())  # aynı TC


def test_raporlar(db, tmp_path):
    db.personel_kaydet(hemsire())
    db.personel_kaydet(Personel(tc="32345678996", ad="Zeynep", soyad="Öztürk", birim="Doğumhane", statu="4B",
                                sozlesme_ucreti=D(40000)))
    sonuclar = db.donem_hesapla("2026-07")
    for ad, fonk in [("kisi.pdf", raporlar.kisi_bordrosu_pdf), ("icmal.pdf", raporlar.icmal_pdf),
                     ("toplu.pdf", raporlar.toplu_bordro_pdf)]:
        yol = tmp_path / ad
        fonk(sonuclar, str(yol), "T.C. SAĞLIK BAKANLIĞI", "Test Hastanesi")
        assert yol.read_bytes()[:4] == b"%PDF"

    raporlar.toplu_bordro_excel(sonuclar, str(tmp_path / "t.xlsx"), "Kurum")
    wb = load_workbook(tmp_path / "t.xlsx")
    assert wb.sheetnames == ["Toplu Bordro", "İcmal"]
    ws = wb["Toplu Bordro"]
    basliklar = [c.value for c in ws[3]]
    net_sutun = basliklar.index("NET") + 1
    assert sum(D(ws.cell(row=r, column=net_sutun).value) for r in (4, 5)) == \
        sum((s.net for s in sonuclar), Decimal(0)).quantize(D("0.01"))

    gruplar = raporlar.icmal(sonuclar)
    assert list(gruplar)[-1] == "GENEL TOPLAM" and gruplar["GENEL TOPLAM"]["kisi"] == 2
    assert gruplar["GENEL TOPLAM"]["net"] == sum((s.net for s in sonuclar), Decimal(0))

    raporlar.banka_listesi_csv(sonuclar, str(tmp_path / "b.csv"))
    with open(tmp_path / "b.csv", encoding="utf-8-sig") as f:
        satirlar = list(csv.reader(f, delimiter=";"))
    assert satirlar[0][0] == "TCKN" and len(satirlar) == 3
    raporlar.banka_listesi_excel(sonuclar, str(tmp_path / "b.xlsx"), "Kurum")
    assert raporlar.eksik_iban(sonuclar) == ["Zeynep Öztürk"]
    assert raporlar.tl(D("1234567.8")) == "1.234.567,80"


def test_excel_iceaktarma(tmp_path):
    yol = str(tmp_path / "sablon.xlsx")
    iceaktar.sablon_olustur(yol)
    liste = iceaktar.personel_oku(yol)
    assert len(liste) == 1
    p = liste[0]
    assert p.tc == "12345678950" and p.derece == 5 and p.es_yardimi is True
    assert p.ozel_hizmet_orani == 100


def test_turkce_sayi_yazimi():
    assert D("3.500,75") == D("3500.75")
    assert D("1,5") == D("1.5")
    assert D("1.575512") == Decimal("1.575512")
    assert D(" 42 000,50 ") == D("42000.50")
    assert D("1.250.000") == D("1250000")
