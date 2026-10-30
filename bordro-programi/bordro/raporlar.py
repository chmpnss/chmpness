"""PDF ve Excel çıktıları: kişi bordrosu, toplu bordro, icmal ve banka listesi."""
from __future__ import annotations

import csv
import os
from collections import OrderedDict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from .hesap import SIFIR, Sonuc

AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos",
         "Eylül", "Ekim", "Kasım", "Aralık"]

# Türkçe karakterler için TTF yazı tipi gerekir (Helvetica ş/ğ/ı içermez).
_FONT_ADAYLARI = [
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]
_FONT = None


def _font() -> tuple[str, str]:
    global _FONT
    if _FONT:
        return _FONT
    for normal, kalin in _FONT_ADAYLARI:
        if os.path.exists(normal):
            pdfmetrics.registerFont(TTFont("BordroFont", normal))
            pdfmetrics.registerFont(TTFont("BordroFont-Bold", kalin if os.path.exists(kalin) else normal))
            _FONT = ("BordroFont", "BordroFont-Bold")
            return _FONT
    raise RuntimeError("Türkçe karakter destekli yazı tipi (Arial/DejaVu Sans) bulunamadı.")


def tl(tutar: Decimal) -> str:
    """1234567.8 -> '1.234.567,80'"""
    return f"{tutar:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def donem_adi(donem: str) -> str:
    yil, ay = donem.split("-")
    return f"{AYLAR[int(ay) - 1]} {yil}"


def _stiller():
    normal, kalin = _font()
    return {
        "baslik": ParagraphStyle("baslik", fontName=kalin, fontSize=12, alignment=1, spaceAfter=2),
        "alt": ParagraphStyle("alt", fontName=normal, fontSize=9, alignment=1),
        "normal": ParagraphStyle("normal", fontName=normal, fontSize=8),
    }


def _tablo_stili(kalin_satirlar=(), baslik=True):
    normal, kalin = _font()
    stil = [("FONTNAME", (0, 0), (-1, -1), normal), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    if baslik:
        stil += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dde6f0")),
                 ("FONTNAME", (0, 0), (-1, 0), kalin)]
    for r in kalin_satirlar:
        stil += [("FONTNAME", (0, r), (-1, r), kalin), ("BACKGROUND", (0, r), (-1, r), colors.HexColor("#f2f2f2"))]
    return TableStyle(stil)


# ---------------------------------------------------------------- Kişi bordrosu
def _kisi_bordrosu_ogeleri(s: Sonuc, kurum: str, birim: str, st) -> list:
    p = s.personel
    oge = [Paragraph(kurum, st["baslik"])]
    if birim:
        oge.append(Paragraph(birim, st["alt"]))
    oge += [Paragraph(f"{donem_adi(s.donem)} Dönemi Maaş Bordrosu", st["alt"]), Spacer(1, 4 * mm)]

    statu = "657 Kadrolu Memur" if p.statu == "657" else "4/B Sözleşmeli Personel"
    kimlik = [
        ["T.C. Kimlik No", p.tc, "Ad Soyad", p.ad_soyad],
        ["Unvan", p.unvan, "Birim", p.birim],
        ["Statü", statu, "IBAN", p.iban],
    ]
    if p.statu == "657":
        kimlik.append(["Derece / Kademe", f"{p.derece} / {p.kademe}", "Ek Gösterge", str(p.ek_gosterge)])
        kimlik.append(["Hizmet Yılı", str(p.hizmet_yili),
                       "Emeklilik", "5434 (Emekli Sandığı)" if p.emeklilik == "5434" else "5510 (SGK 4/c)"])
    t = Table(kimlik, colWidths=[30 * mm, 55 * mm, 25 * mm, 70 * mm])
    t.setStyle(_tablo_stili(baslik=False))
    oge += [t, Spacer(1, 4 * mm)]

    satir_sayisi = max(len(s.odemeler), len(s.kesintiler))
    govde = [["Ödemeler", "Tutar", "Kesintiler", "Tutar"]]
    for i in range(satir_sayisi):
        o = s.odemeler[i] if i < len(s.odemeler) else None
        k = s.kesintiler[i] if i < len(s.kesintiler) else None
        govde.append([o.ad if o else "", tl(o.tutar) if o else "", k.ad if k else "", tl(k.tutar) if k else ""])
    govde.append(["Brüt Toplam", tl(s.brut), "Kesinti Toplamı", tl(s.kesinti_toplami)])
    govde.append(["", "", "NET ÖDENECEK", tl(s.net)])
    t = Table(govde, colWidths=[55 * mm, 30 * mm, 65 * mm, 30 * mm])
    stil = _tablo_stili(kalin_satirlar=(len(govde) - 2, len(govde) - 1))
    stil.add("ALIGN", (1, 0), (1, -1), "RIGHT")
    stil.add("ALIGN", (3, 0), (3, -1), "RIGHT")
    t.setStyle(stil)
    oge += [t, Spacer(1, 4 * mm)]

    bilgi = [["Gelir Vergisi Matrahı", tl(s.gv_matrahi), "Kümülatif GV Matrahı", tl(s.kumulatif_matrah)],
             ["SGK / GSS Prim Matrahı", tl(s.sgk_matrahi), "", ""]]
    for x in s.isveren:
        bilgi.append([x.ad, tl(x.tutar), "", ""])
    t = Table(bilgi, colWidths=[55 * mm, 30 * mm, 65 * mm, 30 * mm])
    stil = _tablo_stili(baslik=False)
    stil.add("ALIGN", (1, 0), (1, -1), "RIGHT")
    stil.add("ALIGN", (3, 0), (3, -1), "RIGHT")
    t.setStyle(stil)
    oge.append(t)
    return oge


def kisi_bordrosu_pdf(sonuclar: list[Sonuc], dosya: str, kurum: str, birim: str = ""):
    """Her personel ayrı sayfada olacak şekilde kişi bordrolarını yazar."""
    st = _stiller()
    doc = SimpleDocTemplate(dosya, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title="Maaş Bordrosu")
    ogeler = []
    for i, s in enumerate(sonuclar):
        if i:
            ogeler.append(PageBreak())
        ogeler += _kisi_bordrosu_ogeleri(s, kurum, birim, st)
    doc.build(ogeler)


# ---------------------------------------------------------------- Toplu / icmal
def _sutunlar(sonuclar: list[Sonuc]) -> tuple[list[str], list[str], list[str]]:
    """Tüm bordrolarda geçen ödeme, kesinti ve işveren satır adları (ilk görülme sırasıyla)."""
    odeme, kesinti, isveren = OrderedDict(), OrderedDict(), OrderedDict()
    for s in sonuclar:
        for x in s.odemeler:
            odeme[x.ad] = None
        for x in s.kesintiler:
            kesinti[x.ad] = None
        for x in s.isveren:
            isveren[x.ad] = None
    return list(odeme), list(kesinti), list(isveren)


def icmal(sonuclar: list[Sonuc], grupla: str | None = "birim") -> "OrderedDict[str, dict]":
    """Kalem bazında toplamlar. grupla='birim' ise birim birim, None ise yalnız genel toplam."""
    gruplar: OrderedDict[str, dict] = OrderedDict()
    for s in sonuclar:
        anahtarlar = ["GENEL TOPLAM"] + ([s.personel.birim or "(Birimsiz)"] if grupla else [])
        for a in anahtarlar:
            g = gruplar.setdefault(a, {"kisi": 0, "odeme": OrderedDict(), "kesinti": OrderedDict(),
                                      "isveren": OrderedDict(), "brut": SIFIR, "kesinti_toplami": SIFIR,
                                      "net": SIFIR})
            g["kisi"] += 1
            for tur, liste in (("odeme", s.odemeler), ("kesinti", s.kesintiler), ("isveren", s.isveren)):
                for x in liste:
                    g[tur][x.ad] = g[tur].get(x.ad, SIFIR) + x.tutar
            g["brut"] += s.brut
            g["kesinti_toplami"] += s.kesinti_toplami
            g["net"] += s.net
    genel = gruplar.pop("GENEL TOPLAM", None)
    if genel:
        gruplar["GENEL TOPLAM"] = genel
    return gruplar


def icmal_pdf(sonuclar: list[Sonuc], dosya: str, kurum: str, birim: str = ""):
    st = _stiller()
    doc = SimpleDocTemplate(dosya, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm, title="Bordro İcmali")
    donem = sonuclar[0].donem if sonuclar else ""
    ogeler = []
    for i, (ad, g) in enumerate(icmal(sonuclar).items()):
        if i:
            ogeler.append(PageBreak())
        ogeler += [Paragraph(kurum, st["baslik"])]
        if birim:
            ogeler.append(Paragraph(birim, st["alt"]))
        ogeler += [Paragraph(f"{donem_adi(donem)} Dönemi Bordro İcmali — {ad} ({g['kisi']} kişi)", st["alt"]),
                   Spacer(1, 4 * mm)]
        satirlar = [["Ödemeler", "Tutar"]] + [[k, tl(v)] for k, v in g["odeme"].items()]
        satirlar.append(["Brüt Toplam", tl(g["brut"])])
        kalin = [len(satirlar) - 1]
        satirlar.append(["Kesintiler", ""])
        kalin.append(len(satirlar) - 1)
        satirlar += [[k, tl(v)] for k, v in g["kesinti"].items()]
        satirlar.append(["Kesinti Toplamı", tl(g["kesinti_toplami"])])
        kalin.append(len(satirlar) - 1)
        satirlar.append(["NET ÖDENECEK", tl(g["net"])])
        kalin.append(len(satirlar) - 1)
        if g["isveren"]:
            satirlar.append(["Kurum (İşveren) Payları", ""])
            kalin.append(len(satirlar) - 1)
            satirlar += [[k, tl(v)] for k, v in g["isveren"].items()]
            satirlar.append(["Kurum Payları Toplamı", tl(sum(g["isveren"].values(), SIFIR))])
            kalin.append(len(satirlar) - 1)
        t = Table(satirlar, colWidths=[110 * mm, 50 * mm])
        stil = _tablo_stili(kalin_satirlar=kalin)
        stil.add("ALIGN", (1, 0), (1, -1), "RIGHT")
        t.setStyle(stil)
        ogeler.append(t)
    doc.build(ogeler)


def toplu_bordro_pdf(sonuclar: list[Sonuc], dosya: str, kurum: str, birim: str = ""):
    """Yatay A4'te personel satırları: temel ödeme/kesinti toplamları."""
    st = _stiller()
    doc = SimpleDocTemplate(dosya, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="Toplu Bordro")
    donem = sonuclar[0].donem if sonuclar else ""
    _, kesinti_adlari, _ = _sutunlar(sonuclar)
    basliklar = ["S.No", "T.C. No", "Ad Soyad", "Unvan", "Brüt"] + kesinti_adlari + ["Kesinti T.", "Net"]
    veri = [basliklar]
    for i, s in enumerate(sonuclar, 1):
        veri.append([str(i), s.personel.tc, s.personel.ad_soyad, s.personel.unvan, tl(s.brut)]
                    + [tl(s.kesinti(a)) for a in kesinti_adlari] + [tl(s.kesinti_toplami), tl(s.net)])
    veri.append(["", "", "TOPLAM", f"{len(sonuclar)} kişi", tl(sum((s.brut for s in sonuclar), SIFIR))]
                + [tl(sum((s.kesinti(a) for s in sonuclar), SIFIR)) for a in kesinti_adlari]
                + [tl(sum((s.kesinti_toplami for s in sonuclar), SIFIR)),
                   tl(sum((s.net for s in sonuclar), SIFIR))])
    baslik_stili = ParagraphStyle("b", fontName=_font()[1], fontSize=6.5, leading=7.5)
    veri[0] = [Paragraph(h, baslik_stili) for h in basliklar]
    t = Table(veri, repeatRows=1)
    stil = _tablo_stili(kalin_satirlar=(len(veri) - 1,))
    stil.add("FONTSIZE", (0, 1), (-1, -1), 7)
    stil.add("ALIGN", (4, 1), (-1, -1), "RIGHT")
    t.setStyle(stil)
    ogeler = [Paragraph(kurum, st["baslik"])]
    if birim:
        ogeler.append(Paragraph(birim, st["alt"]))
    ogeler += [Paragraph(f"{donem_adi(donem)} Dönemi Toplu Maaş Bordrosu", st["alt"]), Spacer(1, 3 * mm), t]
    doc.build(ogeler)


# ---------------------------------------------------------------- Excel
_KALIN = Font(bold=True)
_BASLIK_DOLGU = PatternFill("solid", fgColor="DDE6F0")
_INCE = Side(style="thin", color="999999")
_KENAR = Border(left=_INCE, right=_INCE, top=_INCE, bottom=_INCE)
_PARA = '#,##0.00'


def _baslik_yaz(ws, satir: int, basliklar: list[str]):
    for c, h in enumerate(basliklar, 1):
        h_ = ws.cell(row=satir, column=c, value=h)
        h_.font, h_.fill, h_.border = _KALIN, _BASLIK_DOLGU, _KENAR
        h_.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _genislik(ws, genislikler: dict[int, int]):
    for c, w in genislikler.items():
        ws.column_dimensions[get_column_letter(c)].width = w


def toplu_bordro_excel(sonuclar: list[Sonuc], dosya: str, kurum: str):
    """Tüm kalemler sütun sütun; ayrıca icmal sayfası."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Toplu Bordro"
    donem = sonuclar[0].donem if sonuclar else ""
    ws["A1"] = f"{kurum} — {donem_adi(donem) if donem else ''} Toplu Maaş Bordrosu"
    ws["A1"].font = Font(bold=True, size=12)

    odeme, kesinti, isveren = _sutunlar(sonuclar)
    sabit = ["T.C. No", "Ad", "Soyad", "Unvan", "Birim", "Statü", "Derece", "Kademe", "IBAN"]
    basliklar = (sabit + odeme + ["BRÜT TOPLAM", "GV Matrahı", "Kümülatif GV Matrahı", "SGK Matrahı"]
                 + kesinti + ["KESİNTİ TOPLAMI", "NET"] + isveren)
    _baslik_yaz(ws, 3, basliklar)
    for r, s in enumerate(sonuclar, 4):
        p = s.personel
        degerler = ([p.tc, p.ad, p.soyad, p.unvan, p.birim, p.statu,
                     p.derece if p.statu == "657" else "", p.kademe if p.statu == "657" else "", p.iban]
                    + [s.odeme(a) for a in odeme]
                    + [s.brut, s.gv_matrahi, s.kumulatif_matrah, s.sgk_matrahi]
                    + [s.kesinti(a) for a in kesinti] + [s.kesinti_toplami, s.net]
                    + [sum((x.tutar for x in s.isveren if x.ad == a), SIFIR) for a in isveren])
        for c, v in enumerate(degerler, 1):
            h = ws.cell(row=r, column=c, value=float(v) if isinstance(v, Decimal) else v)
            h.border = _KENAR
            if isinstance(v, Decimal):
                h.number_format = _PARA
    son = 3 + len(sonuclar)
    toplam_satiri = son + 1
    ws.cell(row=toplam_satiri, column=1, value="TOPLAM").font = _KALIN
    for c in range(len(sabit) + 1, len(basliklar) + 1):
        harf = get_column_letter(c)
        h = ws.cell(row=toplam_satiri, column=c, value=f"=SUM({harf}4:{harf}{son})" if sonuclar else 0)
        h.font, h.number_format, h.border = _KALIN, _PARA, _KENAR
    ws.freeze_panes = "D4"
    _genislik(ws, {1: 13, 2: 14, 3: 14, 4: 18, 5: 18, 6: 7, 7: 7, 8: 7, 9: 29})
    for c in range(len(sabit) + 1, len(basliklar) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.row_dimensions[3].height = 45

    # İcmal sayfası
    wi = wb.create_sheet("İcmal")
    gruplar = icmal(sonuclar)
    wi["A1"] = f"{kurum} — {donem_adi(donem) if donem else ''} Bordro İcmali"
    wi["A1"].font = Font(bold=True, size=12)
    grup_adlari = list(gruplar)
    _baslik_yaz(wi, 3, ["Kalem"] + grup_adlari)
    satir = 4

    def yaz(ad, degerler, kalin=False):
        nonlocal satir
        wi.cell(row=satir, column=1, value=ad).border = _KENAR
        if kalin:
            wi.cell(row=satir, column=1).font = _KALIN
        for c, v in enumerate(degerler, 2):
            h = wi.cell(row=satir, column=c, value=float(v) if isinstance(v, Decimal) else v)
            h.border = _KENAR
            if isinstance(v, Decimal):
                h.number_format = _PARA
            if kalin:
                h.font = _KALIN
        satir += 1

    yaz("Personel Sayısı", [g["kisi"] for g in gruplar.values()], kalin=True)
    for a in odeme:
        yaz(a, [g["odeme"].get(a, SIFIR) for g in gruplar.values()])
    yaz("BRÜT TOPLAM", [g["brut"] for g in gruplar.values()], kalin=True)
    for a in kesinti:
        yaz(a, [g["kesinti"].get(a, SIFIR) for g in gruplar.values()])
    yaz("KESİNTİ TOPLAMI", [g["kesinti_toplami"] for g in gruplar.values()], kalin=True)
    yaz("NET ÖDENECEK", [g["net"] for g in gruplar.values()], kalin=True)
    for a in isveren:
        yaz(a, [g["isveren"].get(a, SIFIR) for g in gruplar.values()])
    _genislik(wi, {1: 34})
    for c in range(2, len(grup_adlari) + 2):
        wi.column_dimensions[get_column_letter(c)].width = 18
    wb.save(dosya)


# ---------------------------------------------------------------- Banka listesi
def banka_listesi_excel(sonuclar: list[Sonuc], dosya: str, kurum: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Banka Listesi"
    donem = sonuclar[0].donem if sonuclar else ""
    ws["A1"] = f"{kurum} — {donem_adi(donem) if donem else ''} Maaş Ödeme Listesi"
    ws["A1"].font = Font(bold=True, size=12)
    _baslik_yaz(ws, 3, ["S.No", "T.C. Kimlik No", "Ad Soyad", "IBAN", "Net Tutar", "Açıklama"])
    r = 4
    for i, s in enumerate(sonuclar, 1):
        for c, v in enumerate([i, s.personel.tc, s.personel.ad_soyad, s.personel.iban, float(s.net),
                               f"{donem_adi(donem)} maaş"], 1):
            h = ws.cell(row=r, column=c, value=v)
            h.border = _KENAR
        ws.cell(row=r, column=5).number_format = _PARA
        r += 1
    ws.cell(row=r, column=3, value="TOPLAM").font = _KALIN
    h = ws.cell(row=r, column=5, value=float(sum((s.net for s in sonuclar), SIFIR)))
    h.font, h.number_format = _KALIN, _PARA
    _genislik(ws, {1: 6, 2: 15, 3: 28, 4: 32, 5: 16, 6: 22})
    wb.save(dosya)


def banka_listesi_csv(sonuclar: list[Sonuc], dosya: str):
    """Noktalı virgülle ayrılmış, Türkçe Excel'de doğrudan açılan CSV (ondalık ayırıcı virgül)."""
    with open(dosya, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["TCKN", "AD SOYAD", "IBAN", "TUTAR", "ACIKLAMA"])
        for s in sonuclar:
            w.writerow([s.personel.tc, s.personel.ad_soyad, s.personel.iban,
                        f"{s.net:.2f}".replace(".", ","), f"{donem_adi(s.donem)} maaş"])


def eksik_iban(sonuclar: list[Sonuc]) -> list[str]:
    return [s.personel.ad_soyad for s in sonuclar if not s.personel.iban]
