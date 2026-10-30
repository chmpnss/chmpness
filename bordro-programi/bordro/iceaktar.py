"""Excel'den toplu personel aktarımı ve boş şablon üretimi."""
from __future__ import annotations

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .hesap import D, Personel

# Excel sütun başlığı -> Personel alanı
SUTUNLAR = [
    ("T.C. Kimlik No", "tc"), ("Ad", "ad"), ("Soyad", "soyad"), ("Unvan", "unvan"), ("Birim", "birim"),
    ("Statü (657/4B)", "statu"), ("IBAN", "iban"), ("Derece", "derece"), ("Kademe", "kademe"),
    ("Ek Gösterge", "ek_gosterge"), ("Hizmet Yılı", "hizmet_yili"), ("Yan Ödeme Puanı", "yan_odeme_puani"),
    ("Özel Hizmet Oranı (%)", "ozel_hizmet_orani"), ("Ek Ödeme Oranı (%)", "ek_odeme_orani"),
    ("Sözleşme Ücreti (4B)", "sozlesme_ucreti"), ("Emeklilik (5510/5434)", "emeklilik"),
    ("Eş Yardımı (E/H)", "es_yardimi"), ("Çocuk 0-6 Yaş", "cocuk_0_6"), ("Çocuk 6 Yaş Üstü", "cocuk_6_ustu"),
    ("Engellilik Derecesi (0-3)", "engellilik"), ("Sendika Aidatı", "sendika_aidati"),
    ("Devreden GV Matrahı", "devreden_matrah"),
]
TAMSAYI = {"derece", "kademe", "ek_gosterge", "hizmet_yili", "yan_odeme_puani", "cocuk_0_6",
           "cocuk_6_ustu", "engellilik"}
ONDALIK = {"ozel_hizmet_orani", "ek_odeme_orani", "sozlesme_ucreti", "sendika_aidati", "devreden_matrah"}


def sablon_olustur(dosya: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Personel"
    for c, (baslik, _) in enumerate(SUTUNLAR, 1):
        h = ws.cell(row=1, column=c, value=baslik)
        h.font = Font(bold=True)
        h.fill = PatternFill("solid", fgColor="DDE6F0")
        ws.column_dimensions[h.column_letter].width = max(12, len(baslik) + 2)
    ws.append(["12345678950", "Ayşe", "Yılmaz", "Hemşire", "Acil Servis", "657", "", 5, 3, 0, 12, 2200,
               100, 87, 0, "5510", "E", 1, 1, 0, 250, 0])
    wb.save(dosya)


def _deger(alan, v):
    if v is None or v == "":
        return None
    if alan in TAMSAYI:
        return int(float(str(v).replace(",", ".")))
    if alan in ONDALIK:
        return D(v)
    if alan == "es_yardimi":
        return str(v).strip().upper() in ("E", "EVET", "1", "TRUE", "X")
    v = str(v).strip()
    if alan == "tc" and v.endswith(".0"):
        v = v[:-2]
    return v


def personel_oku(dosya: str) -> list[Personel]:
    ws = load_workbook(dosya, data_only=True).active
    basliklar = [str(h.value or "").strip() for h in ws[1]]
    eslesme = {}
    for baslik, alan in SUTUNLAR:
        if baslik in basliklar:
            eslesme[basliklar.index(baslik)] = alan
    if "tc" not in eslesme.values():
        raise ValueError("Excel dosyasında 'T.C. Kimlik No' sütunu bulunamadı. Şablonu kullanın.")
    sonuc = []
    for satir in ws.iter_rows(min_row=2, values_only=True):
        if not any(satir):
            continue
        veri = {}
        for i, alan in eslesme.items():
            v = _deger(alan, satir[i] if i < len(satir) else None)
            if v is not None:
                veri[alan] = v
        sonuc.append(Personel(**veri))
    return sonuc
